#!/usr/bin/env python3
"""Убирает лист «Обозначения столbцов»; на «Основа»/«Доп» сверху только названия колонок."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

PATH = Path(r"c:\Users\Administrator\Downloads\База Главснаб — итог с обозначениями.xlsx")


def is_block_row(ws) -> bool:
    text = str(ws.cell(1, 1).value or "")
    return text.startswith("Блок 1")


def cleanup_sheet(ws) -> None:
    if is_block_row(ws):
        ws.delete_rows(1)
    # убрать подсказку справа от № п/п
    for col in range(ws.max_column, 12, -1):
        val = ws.cell(2 if ws.cell(1, 2).value == "ИНН" else 1, col).value
        if isinstance(val, str) and val.startswith("Подсказка:"):
            ws.delete_cols(col, ws.max_column - col + 1)
            break
    header_row = 1 if ws.cell(1, 2).value == "ИНН" else 2
    # № п/п оставляем, если уже есть
    last = ws.max_column
    if ws.cell(header_row, last).value != "№ п/п":
        num_col = last + 1
        ws.cell(header_row, num_col, "№ п/п")
        counter = 0
        for row in range(header_row + 1, ws.max_row + 1):
            inn = ws.cell(row, 2).value
            digits = "".join(ch for ch in str(inn or "") if ch.isdigit())
            if len(digits) in (10, 12):
                counter += 1
                ws.cell(row, num_col, counter)
        last = num_col
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(last)}{ws.max_row}"


def main() -> None:
    wb = load_workbook(PATH)
    legend = "Обозначения столбцов"
    if legend in wb.sheetnames:
        del wb[legend]
    for name in wb.sheetnames:
        if name in ("Правила скоринга",):
            continue
        cleanup_sheet(wb[name])
    wb.save(PATH)
    wb.close()
    print(PATH)


if __name__ == "__main__":
    main()
