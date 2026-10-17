#!/usr/bin/env python3
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts.score_glavsnab import RESULT_HEADERS  # noqa: E402

PATH = Path(r"c:\Users\Administrator\Downloads\База Главснаб — итог с обозначениями.xlsx")

BASE = [
    "Организация",
    "ИНН",
    "Основной вид деятельности по ОКВЭД",
    "Телефон",
    "Электронная почта",
    "Основное контактное лицо",
    "Тип отношений",
    "Адрес",
    "Комментарий",
    "Менеджер (Фамилия)",
    "Дата контакта",
    "Результат обзвона",
]
HEADERS = BASE + RESULT_HEADERS + ["№ п/п"]
FILL = PatternFill("solid", fgColor="1F4E79")
FONT = Font(color="FFFFFF", bold=True)
ALIGN = Alignment(wrap_text=True, vertical="center", horizontal="center")


def apply_headers(ws) -> None:
    for col, title in enumerate(HEADERS, 1):
        cell = ws.cell(1, col, title)
        cell.value = title
        cell.fill = FILL
        cell.font = FONT
        cell.alignment = ALIGN
    ws.row_dimensions[1].hidden = False
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"
    ws.sheet_view.topLeftCell = "A1"
    ws.sheet_view.zoomScale = 100


def main() -> None:
    wb = load_workbook(PATH)
    if "Обозначения столбцов" in wb.sheetnames:
        del wb["Обозначения столбцов"]
    for name in wb.sheetnames:
        if name == "Правила скоринга":
            continue
        apply_headers(wb[name])
    wb.active = wb.sheetnames.index("Основа")
    wb.save(PATH)
    wb.close()
    print(PATH)


if __name__ == "__main__":
    main()
