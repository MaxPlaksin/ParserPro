#!/usr/bin/env python3
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

PATH = Path(r"c:\Users\Administrator\Downloads\База Главснаб — итог с обозначениями.xlsx")

LINK_HEADERS = {
    "Ссылка Checko",
    "Ссылка на отчётность",
    "Ссылка на суды",
    "Ссылка на исполнительные",
    "Ссылка на контакты",
    "Ссылка на руководителя",
    "Ссылка на карту",
    "Сайт",
}

LINK_FONT = Font(color="0563C1", underline="single")


def linkify_sheet(ws) -> int:
    headers = {ws.cell(1, col).value: col for col in range(1, ws.max_column + 1)}
    cols = [headers[name] for name in LINK_HEADERS if name in headers]
    count = 0
    for row in range(2, ws.max_row + 1):
        for col in cols:
            value = ws.cell(row, col).value
            if not isinstance(value, str):
                continue
            url = value.strip()
            if not url.startswith(("http://", "https://")):
                continue
            cell = ws.cell(row, col)
            cell.hyperlink = url
            cell.font = LINK_FONT
            count += 1
    return count


def main() -> None:
    wb = load_workbook(PATH)
    total = 0
    for name in wb.sheetnames:
        if name == "Правила скоринга":
            continue
        total += linkify_sheet(wb[name])
    wb.save(PATH)
    wb.close()
    print(PATH)
    print("links", total)


if __name__ == "__main__":
    main()
