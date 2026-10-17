#!/usr/bin/python3
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts.score_glavsnab import RESULT_HEADERS  # noqa: E402

PATH = Path(r"c:\Users\Administrator\Downloads\База Главснаб — итог с обозначениями.xlsx")

BASE = [
    "Организация",
    "ИНН",
    "Основной вид деятельности по ОКВЭД",
    "Телефон",
    "Электронная почта",
    "Основное контактное лицо",
    "Тип отношений",
    "Адрес",
    "Комментарий",
    "Менеджер (Фамилия)",
    "Дата контакта",
    "Результат обзвона",
]
HEADERS = BASE + RESULT_HEADERS + ["№ п/п"]
FILL = PatternFill("solid", fgColor="1F4E79")
FONT = Font(color="FFFFFF", bold=True)


def norm_inn(value) -> bool:
    text = str(value or "").strip()
    if text.endswith(".0"):
        text = text[:-2]
    digits = "".join(ch for ch in text if ch.isdigit())
    return len(digits) in (10, 12)


def repair_sheet(ws) -> None:
    for merged in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged))
    data_start = None
    for row in range(1, ws.max_row + 1):
        val = ws.cell(row, 2).value
        if isinstance(val, str) and val.strip().upper() == "ИНН":
            data_start = row + 1
            break
        if norm_inn(val):
            data_start = row
            break
    if data_start is None:
        data_start = 2
    if data_start > 1:
        ws.delete_rows(1, data_start - 1)
    ws.insert_rows(1)
    for col in range(ws.max_column, len(HEADERS), -1):
        if col > len(HEADERS):
            ws.delete_cols(col)
    for col, title in enumerate(HEADERS, 1):
        cell = ws.cell(1, col, title)
        cell.fill = FILL
        cell.font = FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    counter = 0
    num_col = len(HEADERS)
    for row in range(2, ws.max_row + 1):
        if norm_inn(ws.cell(row, 2).value):
            counter += 1
            ws.cell(row, num_col, counter)
        else:
            ws.cell(row, num_col, None)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"


def main() -> None:
    wb = load_workbook(PATH)
    if "Обозначения столбцов" in wb.sheetnames:
        del wb["Обозначения столбцов"]
    for name in wb.sheetnames:
        if name != "Правила скоринга":
            repair_sheet(wb[name])
    wb.save(PATH)
    wb.close()
    print(PATH)


if __name__ == "__main__":
    main()
