#!/usr/bin/env python3
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from scripts.score_glavsnab import RESULT_HEADERS  # noqa: E402

PATH = Path(r"c:\Users\Administrator\Downloads\База Главснаб — итог с обозначениями.xlsx")

BASE_HEADERS = [
    "Организация",
    "ИНН",
    "Основной вид деятельности по ОКВЭД",
    "Телефон",
    "Электронная почта",
    "Основное контактное лицо",
    "Тип отношений",
    "Адрес",
    "Комментарий",
    "Менеджер (Фамилия)",
    "Дата контакта",
    "Результат обзвона",
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def norm_inn(value) -> bool:
    if value is None:
        return False
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    digits = "".join(ch for ch in text if ch.isdigit())
    return len(digits) in (10, 12)


def write_header_row(ws) -> None:
    headers = BASE_HEADERS + RESULT_HEADERS + ["№ п/п"]
    if ws.cell(1, 2).value == "ИНН":
        row = 1
    else:
        ws.insert_rows(1)
        row = 1
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row, col, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(headers))}{ws.max_row}"


def fill_row_numbers(ws) -> None:
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    if "№ п/п" not in headers:
        return
    col = headers.index("№ п/п") + 1
    counter = 0
    for row in range(2, ws.max_row + 1):
        if norm_inn(ws.cell(row, 2).value):
            counter += 1
            ws.cell(row, col, counter)
        else:
            ws.cell(row, col, None)


def main() -> None:
    wb = load_workbook(PATH)
    if "Обозначения столбцов" in wb.sheetnames:
        del wb["Обозначения столбцов"]
    for name in wb.sheetnames:
        if name == "Правила скоринга":
            continue
        ws = wb[name]
        if ws.cell(1, 2).value != "ИНН":
            write_header_row(ws)
        else:
            headers = BASE_HEADERS + RESULT_HEADERS + ["№ п/п"]
            for col, header in enumerate(headers, 1):
                ws.cell(1, col, header)
                cell = ws.cell(1, col)
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
        fill_row_numbers(ws)
    wb.save(PATH)
    wb.close()
    print(PATH)


if __name__ == "__main__":
    main()
