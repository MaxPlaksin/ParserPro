#!/usr/bin/env python3
"""Прогон базы Главснаб через Checko и светофор надёжности.

Исходные колонки не перезаписываются: результат дописывается справа.
Повторный запуск продолжает с уже сохранённых карточек.
"""

from __future__ import annotations

import argparse
import asyncio
import logging
import os
import re
import shutil
import sys
import time
from dataclasses import asdict
from datetime import datetime
from html import unescape
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from common import CloakSession, atomic_write_json, setup_logging
from counterparty.providers.checko_web import CheckoWebProvider
from counterparty.version import PARSER_VERSION
from normalizer.checko import build_company_check_record

logger = logging.getLogger("score_glavsnab")

DEFAULT_XLSX = Path(r"c:\Users\Administrator\Downloads\База Главснаб.xlsx")
CACHE_DIR = ROOT / "output" / "glavsnab_checko" / "cache"

# Пороги из opisanie.txt (внутренний скоринг). «Крупная сумма» в ТЗ не задана числом:
# требования ответчика от 1 млн руб. или непогашенная задолженность по ИП от 500 тыс. руб.
REVENUE_DROP_PCT = -30.0
ACTIVE_ENFORCEMENT_LIMIT = 3
ARBITRATION_DEFENDANT_LIMIT = 5
LARGE_CLAIM_RUB = 1_000_000.0
LARGE_DEBT_RUB = 500_000.0
RISK_POINT = 20

RESULT_HEADERS = [
    "Светофор",
    "Балл риска",
    "Причины",
    "Статус компании",
    "Банкротство",
    "Выручка, год 1",
    "Выручка год 1, руб",
    "Выручка, год 2",
    "Выручка год 2, руб",
    "Динамика выручки, %",
    "Прибыль, год 1",
    "Прибыль год 1, руб",
    "Прибыль, год 2",
    "Прибыль год 2, руб",
    "Динамика прибыли, %",
    "ИП всего",
    "ИП активных",
    "Задолженность по ИП, руб",
    "Арбитраж ответчик, дел",
    "Арбитраж ответчик, активных",
    "Сумма требований ответчика, руб",
    "Оценка Checko, %",
    "Название",
    "ОГРН",
    "КПП",
    "Дата регистрации",
    "ОКВЭД",
    "Юридический адрес",
    "Руководитель",
    "ИНН руководителя",
    "Телефоны",
    "Электронная почта",
    "Сайт",
    "Ссылка Checko",
    "Ссылка на отчётность",
    "Ссылка на суды",
    "Ссылка на исполнительные",
    "Ссылка на контакты",
    "Ссылка на руководителя",
    "Ссылка на карту",
    "Ошибка",
    "Проверено",
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
GREEN_FONT = Font(color="006100", bold=True)
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
RED_FONT = Font(color="9C0006", bold=True)
YELLOW_FILL = PatternFill("solid", fgColor="FFF2CC")
YELLOW_FONT = Font(color="806000", bold=True)
WRAP = Alignment(wrap_text=True, vertical="center")

STATUS_RE = re.compile(
    r'class="text-(success|danger|warning) fw-[67]00">\s*([^<]{3,220})',
    re.I,
)
COMPANY_STATE_RE = re.compile(
    r"действующ|ликвидир|банкрот|прекрат|исключен|реорганизац",
    re.I,
)
RATING_RE = re.compile(
    r'id="rating-bar"[\s\S]{0,800}?class="green"[^>]*>\s*([+\-]?\d+)\s*%',
    re.I,
)
EFRSB_RE = re.compile(r"ЕФРСБ[\s\S]{0,1200}?</div>\s*</div>", re.I)
TAG_RE = re.compile(r"<[^>]+>")


def norm_inn(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if value != value:
            return ""
        value = int(value)
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    digits = "".join(ch for ch in text if ch.isdigit())
    if len(digits) in (10, 12):
        return digits
    return ""


def _section(html: str, section_id: str) -> str:
    match = re.search(
        rf'<section id="{re.escape(section_id)}"[\s\S]*?</section>',
        html or "",
        re.I,
    )
    return match.group(0) if match else ""


def _plain(html: str) -> str:
    text = TAG_RE.sub(" ", html or "")
    return unescape(re.sub(r"\s+", " ", text)).strip()


def parse_status(html: str) -> tuple[str, str]:
    status = ""
    notes: list[str] = []
    for match in STATUS_RE.finditer(html or ""):
        text = re.sub(r"\s+", " ", match.group(2)).strip()
        low = text.lower()
        if not status and COMPANY_STATE_RE.search(text):
            status = text
        if "недостоверн" in low and not any("недостоверн" in note for note in notes):
            notes.append("сведения ЕГРЮЛ недостоверны")
    if notes:
        status = f"{status}; {notes[0]}" if status else notes[0]
    bankruptcy = ""
    block = EFRSB_RE.search(html or "")
    if block:
        plain = _plain(block.group(0))
        if "Нет сообщений о банкротстве" in plain:
            bankruptcy = "нет"
        elif re.search(r"банкрот", plain, re.I):
            bankruptcy = "есть сведения"
    if not bankruptcy and re.search(r"банкротств", status, re.I):
        bankruptcy = "есть сведения"
    return status, bankruptcy


def parse_checko_rating(html: str) -> int | None:
    match = RATING_RE.search(html or "")
    if not match:
        return None
    try:
        return int(match.group(1))
    except ValueError:
        return None


def _inactive(status: str, bankruptcy: str) -> bool:
    low = (status or "").lower()
    if bankruptcy == "есть сведения":
        return True
    markers = (
        "ликвидир",
        "банкрот",
        "прекрат",
        "исключен",
        "недействующ",
    )
    return any(marker in low for marker in markers)


def _label_value(html: str, label: str) -> str:
    match = re.search(
        rf'class="fw-700">{re.escape(label)}</div>\s*<div>([^<]+)',
        html or "",
        re.I,
    )
    if not match:
        return ""
    return unescape(re.sub(r"\s+", " ", match.group(1))).strip()


def _copy_field(html: str, element_id: str) -> str:
    match = re.search(rf'id="{re.escape(element_id)}"[^>]*>([^<]+)', html or "")
    return unescape(match.group(1)).strip() if match else ""


def parse_chart_finances(html: str) -> dict[int, dict[str, float | None]]:
    """Точные рубли из графика на карточке. Страница отчётности для этого не нужна."""
    page = html or ""
    labels = re.search(r"const labels=\[([^\]]+)\]", page)
    years = [int(year) for year in re.findall(r"'(\d{4})'", labels.group(1))] if labels else []

    def series(code: str) -> list[float]:
        match = re.search(rf'code==="{code}"\)series=\[\[(.*?)\]\]', page)
        if not match:
            return []
        return [float(value) for value in re.findall(r"value:'(-?\d+)'", match.group(1))]

    revenue = series("2110")
    profit = series("2400")
    found: dict[int, dict[str, float | None]] = {}
    for index, year in enumerate(years):
        item: dict[str, float | None] = {}
        if index < len(revenue):
            item["revenue"] = revenue[index]
        if index < len(profit):
            item["profit"] = profit[index]
        if item:
            found[year] = item
    return found


def apply_chart_finances(record, html: str) -> None:
    if record.revenue_value_2 is not None or record.profit_value_2 is not None:
        return
    finances = parse_chart_finances(html)
    years = sorted(finances)
    if not years:
        return
    if len(years) == 1:
        year = years[-1]
        earlier = later = year
    else:
        earlier, later = years[-2], years[-1]
    record.revenue_year_1 = earlier
    record.revenue_year_2 = later
    record.revenue_value_1 = finances[earlier].get("revenue")
    record.revenue_value_2 = finances[later].get("revenue")
    record.profit_year_1 = earlier
    record.profit_year_2 = later
    record.profit_value_1 = finances[earlier].get("profit")
    record.profit_value_2 = finances[later].get("profit")
    old = record.revenue_value_1
    new = record.revenue_value_2
    if old not in (None, 0) and new is not None:
        record.revenue_change_percent = round((new - old) / abs(old) * 100, 2)
    old_profit = record.profit_value_1
    new_profit = record.profit_value_2
    if old_profit not in (None, 0) and new_profit is not None:
        record.profit_change_percent = round((new_profit - old_profit) / abs(old_profit) * 100, 2)


def section_links(url: str) -> dict[str, str]:
    clean = (url or "").split("?")[0].strip()
    links = {
        "Ссылка Checko": clean,
        "Ссылка на отчётность": "",
        "Ссылка на суды": "",
        "Ссылка на исполнительные": "",
        "Ссылка на контакты": "",
        "Ссылка на руководителя": "",
    }
    match = OGRN_IN_URL_RE.search(clean)
    if "/company/" in clean and match and len(match.group(1)) == 13:
        base = f"https://checko.ru/company/{match.group(1)}"
        links["Ссылка на отчётность"] = f"{base}/finances/accounting"
        links["Ссылка на суды"] = f"{base}/legal-cases"
        links["Ссылка на исполнительные"] = f"{base}/enforcements"
        links["Ссылка на контакты"] = f"{base}/contacts"
        links["Ссылка на руководителя"] = f"{base}/management"
    return links


def parse_profile(html: str, url: str) -> dict[str, str]:
    """Карточка с уже скачанной страницы: реквизиты, контакты и ссылки."""
    page = html or ""
    name = ""
    name_match = re.search(r"<h1[^>]*>(.*?)</h1>", page, re.I | re.S)
    if name_match:
        name = unescape(re.sub(r"<[^>]+>", "", name_match.group(1)))
        name = re.sub(r"\s+", " ", name).strip()
    okved = ""
    okved_match = re.search(
        r'id="activity-name"[^>]*onclick="text_to_cb\(\'([^\']*)\'[^"]*"[^>]*>([^<]+)',
        page,
    )
    if okved_match:
        okved = f"{unescape(okved_match.group(2)).strip()} {unescape(okved_match.group(1)).strip()}".strip()
    director = ""
    director_inn = ""
    director_url = ""
    director_match = re.search(
        r'class="fw-700">[^<]*директор[^<]*</div>\s*<a class="link" href="([^"]+)">([^<]+)</a>',
        page,
        re.I,
    )
    if director_match:
        director_url = director_match.group(1)
        director = unescape(director_match.group(2)).strip()
        inn_match = re.search(r"/person/(\d+)", director_url)
        if inn_match:
            director_inn = inn_match.group(1)
        if director_url.startswith("/"):
            director_url = "https://checko.ru" + director_url
    contacts = _section(page, "contacts") or ""
    phones: list[str] = []
    for label in re.findall(r'href="tel:[^"]+"[^>]*>([^<]*)', contacts, re.I):
        phone = unescape(label).strip()
        if phone and phone not in phones and phone != "—":
            phones.append(phone)
    emails: list[str] = []
    for mail in re.findall(r'href="mailto:([^"?]+)', contacts, re.I):
        mail = unescape(mail).strip()
        if mail and mail not in emails:
            emails.append(mail)
    site = ""
    site_match = re.search(
        r"Веб-сайт</strong>\s*<a[^>]+href=\"(https?://[^\"]+)\"[^>]*>([^<]*)",
        contacts,
        re.I,
    )
    if site_match:
        site = site_match.group(1)
    map_match = re.search(r'href="(https://yandex\.ru/maps/\?text=[^"]+)"', contacts, re.I)
    profile = {
        "Название": name,
        "ОГРН": _copy_field(page, "copy-ogrn"),
        "КПП": _copy_field(page, "copy-kpp"),
        "Дата регистрации": _label_value(page, "Дата регистрации"),
        "ОКВЭД": okved,
        "Юридический адрес": _copy_field(page, "copy-address") or _copy_field(page, "copy-x-address"),
        "Руководитель": director,
        "ИНН руководителя": director_inn,
        "Телефоны": ", ".join(phones),
        "Электронная почта": ", ".join(emails),
        "Сайт": site,
        "Ссылка на карту": unescape(map_match.group(1)) if map_match else "",
    }
    profile.update(section_links(url))
    if director_url:
        profile["Ссылка на руководителя"] = director_url
    return profile


def complete_profile(payload: dict) -> dict:
    for key, value in section_links(payload.get("Ссылка Checko") or "").items():
        if value and not payload.get(key):
            payload[key] = value
    return payload


def score_record(
    record,
    *,
    status: str,
    bankruptcy: str,
    checko_rating: int | None,
    error: str,
    profile: dict | None = None,
) -> dict:
    reasons: list[str] = []
    points = 0

    def add(reason: str) -> None:
        nonlocal points
        points += RISK_POINT
        reasons.append(reason)

    change = record.revenue_change_percent
    if change is not None and change <= REVENUE_DROP_PCT:
        add(f"падение выручки {change}%")

    profit = record.profit_value_2
    if profit is not None and profit < 0:
        add("убыток")

    active_enf = record.enforcement_active_count
    if active_enf is None:
        active_enf = record.enforcement_total_count
    if active_enf is not None and active_enf > ACTIVE_ENFORCEMENT_LIMIT:
        add(f"активные ИП {active_enf}")

    cases = record.arbitration_defendant_total_count
    if cases is not None and cases > ARBITRATION_DEFENDANT_LIMIT:
        add(f"арбитраж ответчик {cases} дел")

    claim = record.arbitration_defendant_total_amount or 0
    debt = record.enforcement_total_amount or 0
    if claim >= LARGE_CLAIM_RUB:
        add(f"сумма требований {int(claim)} руб.")
    elif debt >= LARGE_DEBT_RUB:
        add(f"задолженность по ИП {int(debt)} руб.")

    if _inactive(status, bankruptcy):
        if "недействующая" not in " ".join(reasons):
            reasons.append(status or "банкротство")
        points = max(points, RISK_POINT)

    has_finance = record.revenue_value_2 is not None or record.profit_value_2 is not None
    if error:
        light = "Нет данных"
    elif points > 0:
        light = "Рискованный"
    elif not has_finance and not status:
        light = "Нет данных"
        reasons.append("нет отчётности и статуса")
    elif not has_finance:
        light = "Нет данных"
        reasons.append("нет финансовой отчётности")
    else:
        light = "Надежный"

    checked = datetime.now().strftime("%Y-%m-%d %H:%M")
    result = {
        "Светофор": light,
        "Балл риска": points if not error else None,
        "Причины": "; ".join(reasons),
        "Статус компании": status,
        "Банкротство": bankruptcy,
        "Выручка, год 1": record.revenue_year_1,
        "Выручка год 1, руб": record.revenue_value_1,
        "Выручка, год 2": record.revenue_year_2,
        "Выручка год 2, руб": record.revenue_value_2,
        "Динамика выручки, %": record.revenue_change_percent,
        "Прибыль, год 1": record.profit_year_1,
        "Прибыль год 1, руб": record.profit_value_1,
        "Прибыль, год 2": record.profit_year_2,
        "Прибыль год 2, руб": record.profit_value_2,
        "Динамика прибыли, %": record.profit_change_percent,
        "ИП всего": record.enforcement_total_count,
        "ИП активных": record.enforcement_active_count,
        "Задолженность по ИП, руб": record.enforcement_total_amount,
        "Арбитраж ответчик, дел": record.arbitration_defendant_total_count,
        "Арбитраж ответчик, активных": record.arbitration_defendant_active_count,
        "Сумма требований ответчика, руб": record.arbitration_defendant_total_amount,
        "Оценка Checko, %": checko_rating,
        "Ошибка": error,
        "Проверено": checked,
        "_name": record.name,
    }
    result.update(section_links(record.url))
    for key, value in (profile or {}).items():
        if value:
            result[key] = value
    if record.name and not result.get("Название"):
        result["Название"] = record.name
    if record.ogrn and not result.get("ОГРН"):
        result["ОГРН"] = record.ogrn
    return result


def empty_result(error: str) -> dict:
    row = {header: None for header in RESULT_HEADERS}
    row["Светофор"] = "Нет данных"
    row["Ошибка"] = error
    row["Причины"] = error
    row["Проверено"] = datetime.now().strftime("%Y-%m-%d %H:%M")
    return row


def cache_path(inn: str) -> Path:
    return CACHE_DIR / f"{inn}.json"


def write_cache(inn: str, payload: dict) -> None:
    path = cache_path(inn)
    if path.exists() and payload.get("Ошибка"):
        import json

        try:
            previous = json.loads(path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            previous = {}
        if previous.get("Светофор") and not previous.get("Ошибка"):
            logger.warning("Успешную карточку %s не затираю ошибкой", inn)
            return
    atomic_write_json(path, payload)


def load_cached(inn: str) -> dict | None:
    path = cache_path(inn)
    if not path.exists():
        return None
    import json

    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return None
    if not payload.get("Светофор"):
        return None
    has_facts = any(
        payload.get(key) not in (None, "")
        for key in (
            "Выручка год 2, руб",
            "Прибыль год 2, руб",
            "Арбитраж ответчик, дел",
            "ИП всего",
            "Банкротство",
            "Статус компании",
        )
    )
    if not has_facts and not (payload.get("Ошибка") or "").strip():
        return None
    error = (payload.get("Ошибка") or "").lower()
    if any(token in error for token in (
        "blocked", "timeout", "timed out", "net::", "target closed", "429", "403",
        "connection closed", "browser has been closed", "target page",
    )):
        return None
    if len(inn) == 12 and "not found" in error:
        return None
    return payload


CARD_HREF_RE = re.compile(r'href="(?:https://checko\.ru)?(/(?:company|entrepreneur)/[^"#?]+)"', re.I)
OGRN_IN_URL_RE = re.compile(r"(?:-|/)(\d{13}|\d{15})(?:/|$)")
BLOCK_BACKOFF_SEC = (180, 180, 300)


def _is_block(result: dict) -> bool:
    error = (result.get("Ошибка") or "").lower()
    return any(token in error for token in ("429", "403", "blocked"))


def _is_dead_browser(result: dict) -> bool:
    error = (result.get("Ошибка") or "").lower()
    return any(token in error for token in ("connection closed", "target closed", "browser has been closed", "target page"))


async def _resolve_card(provider: CheckoWebProvider, page, inn: str) -> tuple[str, str]:
    from counterparty.providers.checko_web import CHECKO_SEARCH

    html, final_url, _ = await provider._fetch_html(page, CHECKO_SEARCH.format(inn=inn))
    if "/company/" in final_url or "/entrepreneur/" in final_url:
        return html, final_url
    match = CARD_HREF_RE.search(html or "")
    if not match:
        raise RuntimeError(f"Company not found for INN {inn}")
    html, final_url, _ = await provider._fetch_html(page, "https://checko.ru" + match.group(1))
    if "/company/" not in final_url and "/entrepreneur/" not in final_url:
        raise RuntimeError(f"Company not found for INN {inn}")
    return html, final_url


async def fetch_one(provider: CheckoWebProvider, page, inn: str) -> dict:
    try:
        main_html, company_url = await _resolve_card(provider, page, inn)
        ogrn_match = OGRN_IN_URL_RE.search(company_url)
        ogrn = ogrn_match.group(1) if ogrn_match else ""
        accounting_html = ""
        if "/company/" in company_url and ogrn:
            accounting_url = f"https://checko.ru/company/{ogrn}/finances/accounting"
            accounting_html, _, _ = await provider._fetch_html(page, accounting_url)
        legal_html = _section(main_html, "legal-cases") or main_html
        enforcement_html = _section(main_html, "enforcements") or main_html
        record = build_company_check_record(
            inn=inn,
            main_html=main_html,
            accounting_html=accounting_html,
            legal_html=legal_html,
            enforcement_html=enforcement_html,
            company_url=company_url,
            source=provider.name,
            parser_version=PARSER_VERSION,
        )
        enforcement = (record.raw_data or {}).get("enforcement") or {}
        debt = enforcement.get("debt_amount")
        if debt is not None:
            record.enforcement_total_amount = debt
        if record.enforcement_total_count is None and record.enforcement_active_count is None:
            plain = _plain(enforcement_html).lower()
            if "нет сведений" in plain or "не найден" in plain or "отсутств" in plain:
                record.enforcement_total_count = 0
                record.enforcement_active_count = 0
        legal_plain = _plain(legal_html).lower()
        if record.arbitration_defendant_total_count is None and (
            "нет сведений" in legal_plain or "не найден" in legal_plain or "отсутств" in legal_plain
        ):
            record.arbitration_defendant_total_count = 0
            record.arbitration_defendant_active_count = 0
        apply_chart_finances(record, main_html)
        status, bankruptcy = parse_status(main_html)
        rating = parse_checko_rating(main_html)
        result = score_record(
            record,
            status=status,
            bankruptcy=bankruptcy,
            checko_rating=rating,
            error="",
            profile=parse_profile(main_html, company_url),
        )
        result["_raw_status"] = "ok"
        return result
    except Exception as exc:
        logger.error("FAIL INN %s: %s", inn, exc)
        result = empty_result(str(exc))
        result["_raw_status"] = "error"
        return result


def _sheet_header_row(ws) -> int:
    value = ws.cell(1, 2).value
    if isinstance(value, str) and value.strip().upper() == "ИНН":
        return 1
    return 0


def _ensure_dop_header(ws) -> None:
    """Лист «Доп» пришёл без шапки. Вставляем строку заголовков один раз."""
    if _sheet_header_row(ws):
        return
    ws.insert_rows(1)
    headers = [
        "Организация",
        "ИНН",
        "Основной вид деятельности по ОКВЭД",
        "Телефон",
        "Электронная почта",
        "Основное контактное лицо",
        "Тип отношений",
        "Адрес",
        "Комментарий",
        "Менеджер (Фамилия)",
        "Дата контакта",
        "Результат обзвона",
    ]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT


def _result_start_col(ws) -> int:
    header_row = _sheet_header_row(ws) or 1
    for col in range(1, ws.max_column + 1):
        if ws.cell(header_row, col).value == RESULT_HEADERS[0]:
            return col
    return ws.max_column + 1


def _write_headers(ws, start_col: int) -> None:
    header_row = _sheet_header_row(ws) or 1
    for offset, header in enumerate(RESULT_HEADERS):
        cell = ws.cell(header_row, start_col + offset, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(start_col + len(RESULT_HEADERS) - 1)}{ws.max_row}"
    ws.freeze_panes = "A2"
    widths = {
        "Светофор": 16,
        "Балл риска": 12,
        "Причины": 42,
        "Статус компании": 28,
        "Ссылка Checko": 42,
        "Название": 36,
        "Юридический адрес": 46,
        "ОКВЭД": 36,
        "Руководитель": 32,
        "Телефоны": 28,
        "Ошибка": 36,
    }
    for offset, header in enumerate(RESULT_HEADERS):
        letter = get_column_letter(start_col + offset)
        ws.column_dimensions[letter].width = widths.get(header, 22)


def _paint(cell, light: str | None) -> None:
    if light == "Надежный":
        cell.fill = GREEN_FILL
        cell.font = GREEN_FONT
    elif light == "Рискованный":
        cell.fill = RED_FILL
        cell.font = RED_FONT
    elif light == "Нет данных":
        cell.fill = YELLOW_FILL
        cell.font = YELLOW_FONT


def apply_results(ws, by_inn: dict[str, dict]) -> int:
    if ws.title == "Доп":
        _ensure_dop_header(ws)
    start_col = _result_start_col(ws)
    header_row = _sheet_header_row(ws) or 1
    _write_headers(ws, start_col)
    written = 0
    for row in range(header_row + 1, ws.max_row + 1):
        inn = norm_inn(ws.cell(row, 2).value)
        if not inn:
            payload = empty_result("ИНН не указан или некорректный")
        else:
            payload = by_inn.get(inn)
            if not payload:
                continue
        for offset, header in enumerate(RESULT_HEADERS):
            cell = ws.cell(row, start_col + offset, payload.get(header))
            if header in ("Причины", "Ошибка", "Ссылка Checko", "Статус компании", "Юридический адрес", "ОКВЭД", "Название"):
                cell.alignment = WRAP
            value = payload.get(header)
            if isinstance(value, str) and value.startswith("http"):
                cell.hyperlink = value
                cell.style = "Hyperlink"
            if header.endswith(", руб") or header == "Балл риска":
                cell.number_format = "#,##0"
            if header.startswith("Динамика"):
                cell.number_format = "0.00"
        light_cell = ws.cell(row, start_col)
        _paint(light_cell, payload.get("Светофор"))
        written += 1
    return written


def write_rules_sheet(wb) -> None:
    if "Правила скоринга" in wb.sheetnames:
        del wb["Правила скоринга"]
    ws = wb.create_sheet("Правила скоринга")
    lines = [
        "Светофор надёжности контрагента",
        "",
        "Надёжный — компания действующая, есть отчётность, ни один порог ниже не сработал.",
        "Рискованный — сработал хотя бы один порог, либо компания ликвидирована / в банкротстве.",
        "Нет данных — Checko не отдал карточку или нет финансовой отчётности. Цифры по судам, если они есть, всё равно записаны.",
        "",
        "Балл риска. Каждый сработавший порог даёт +20. Это тот расчёт, который уже использовался:",
        f"падение выручки сильнее {abs(REVENUE_DROP_PCT):.0f}% за последний доступный год к предыдущему — +20",
        "чистая прибыль последнего года меньше нуля (убыток) — +20",
        f"активных исполнительных производств больше {ACTIVE_ENFORCEMENT_LIMIT} — +20",
        f"арбитражных дел, где компания ответчик, больше {ARBITRATION_DEFENDANT_LIMIT} — +20",
        f"сумма требований к ответчику от {int(LARGE_CLAIM_RUB):,} руб. — +20".replace(",", " "),
        f"если сумма требований меньше порога, но задолженность по ИП от {int(LARGE_DEBT_RUB):,} руб. — +20".replace(",", " "),
        "",
        "Ликвидация, прекращение деятельности и сообщения о банкротстве сами по себе делают контрагента рискованным.",
        "Отметка «сведения ЕГРЮЛ недостоверны» пишется в статус, но сама по себе балл не увеличивает: её видно в колонке и можно отфильтровать вручную.",
        "",
        "Годы выручки и прибыли — два последних года, которые есть в отчётности Checko (коды 2110 и 2400).",
        "Оценка Checko, % — их собственная шкала на карточке. В светофор она не входит, оставлена для ручной сверки.",
    ]
    for idx, line in enumerate(lines, 1):
        cell = ws.cell(idx, 1, line)
        if idx == 1:
            cell.font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 120


def _exclusive_lock(path: Path):
    lock = path.with_name(path.name + ".lock")
    for _ in range(40):
        try:
            return os.open(str(lock), os.O_CREAT | os.O_EXCL | os.O_RDWR), lock
        except FileExistsError:
            time.sleep(0.5)
    try:
        lock.unlink()
    except OSError:
        pass
    return os.open(str(lock), os.O_CREAT | os.O_EXCL | os.O_RDWR), lock


def save_workbook(path: Path, by_inn: dict[str, dict]) -> Path:
    for payload in by_inn.values():
        complete_profile(payload)
    fd, lock = _exclusive_lock(path)
    try:
        return _save_workbook(path, by_inn)
    finally:
        os.close(fd)
        try:
            lock.unlink()
        except OSError:
            pass


def _save_workbook(path: Path, by_inn: dict[str, dict]) -> Path:
    wb = load_workbook(path)
    for name in wb.sheetnames:
        if name == "Правила скоринга":
            continue
        apply_results(wb[name], by_inn)
    write_rules_sheet(wb)
    tmp = path.with_name(path.stem + ".__tmp__.xlsx")
    wb.save(tmp)
    wb.close()
    try:
        tmp.replace(path)
    except PermissionError:
        alt = path.with_name(path.stem + " — скоринг.xlsx")
        tmp.replace(alt)
        logger.error("Исходный файл занят. Результат записан в %s", alt)
        return alt
    return path


def collect_inns(path: Path) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    found: list[str] = []
    seen: set[str] = set()
    for name in wb.sheetnames:
        if name == "Правила скоринга":
            continue
        ws = wb[name]
        rows = ws.iter_rows(min_row=1, max_col=2, values_only=True)
        for idx, row in enumerate(rows):
            org, inn_value = (row + (None, None))[:2]
            if idx == 0 and isinstance(inn_value, str) and inn_value.strip().upper() == "ИНН":
                continue
            inn = norm_inn(inn_value)
            if inn and inn not in seen:
                seen.add(inn)
                found.append(inn)
            elif org and not inn:
                logger.info("Пропуск без ИНН на листе %s: %s", name, org)
    wb.close()
    return found


async def _fetch_until_open(provider: CheckoWebProvider, page, inn: str) -> dict:
    """Повторяет тот же ИНН, пока Checko не снимет ограничение. К следующему не переходит."""
    result = await fetch_one(provider, page, inn)
    attempt = 0
    while _is_block(result):
        wait = BLOCK_BACKOFF_SEC[min(attempt, len(BLOCK_BACKOFF_SEC) - 1)]
        attempt += 1
        logger.warning(
            "Checko ограничил доступ, пауза %s с, попытка %s, ИНН %s",
            wait,
            attempt,
            inn,
        )
        await asyncio.sleep(wait)
        result = await fetch_one(provider, page, inn)
    return result


async def run(path: Path, *, limit: int | None, delay: float, gap: float, headless: bool) -> None:
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    backup = path.with_name(path.stem + " — до скоринга.xlsx")
    if not backup.exists():
        shutil.copy2(path, backup)
        logger.info("Резервная копия: %s", backup)

    inns = collect_inns(path)
    if limit:
        inns = inns[:limit]
    logger.info("ИНН к проверке: %s", len(inns))

    by_inn: dict[str, dict] = {}
    pending: list[str] = []
    for inn in inns:
        cached = load_cached(inn)
        if cached:
            by_inn[inn] = cached
        else:
            pending.append(inn)
    logger.info("Уже в кэше: %s, осталось: %s", len(by_inn), len(pending))
    if by_inn:
        path = save_workbook(path, by_inn)

    if not pending:
        logger.info("Новых ИНН нет, файл обновлён из кэша")
        return

    provider = CheckoWebProvider(
        headless=headless,
        delay_sec=delay,
        profile_dir=str(ROOT / "output" / "glavsnab_checko" / "cloak_profile"),
    )
    done_since_save = 0
    async with CloakSession(headless=headless, humanize=True, profile_dir=provider.profile_dir) as session:
        page = await session.new_page()
        try:
            for idx, inn in enumerate(pending, 1):
                logger.info("[%s/%s] %s", idx, len(pending), inn)
                fresh = load_cached(inn)
                if fresh:
                    by_inn[inn] = fresh
                    logger.info("уже в кэше, пропуск %s %s", inn, fresh.get("Светофор"))
                    continue
                result = await _fetch_until_open(provider, page, inn)
                if _is_dead_browser(result):
                    logger.warning("Браузер оборвался, открываю заново и повторяю %s", inn)
                    try:
                        await page.close()
                    except Exception:
                        pass
                    await asyncio.sleep(5)
                    page = await session.new_page()
                    result = await _fetch_until_open(provider, page, inn)
                public = {key: value for key, value in result.items() if not key.startswith("_")}
                write_cache(inn, public)
                by_inn[inn] = public
                done_since_save += 1
                logger.info(
                    "%s %s score=%s revenue=%s profit=%s enf=%s cases=%s",
                    inn,
                    public.get("Светофор"),
                    public.get("Балл риска"),
                    public.get("Выручка год 2, руб"),
                    public.get("Прибыль год 2, руб"),
                    public.get("ИП активных"),
                    public.get("Арбитраж ответчик, дел"),
                )
                if done_since_save >= 1 or idx == len(pending):
                    path = save_workbook(path, by_inn)
                    done_since_save = 0
                    logger.info("Файл обновлён, готово %s из %s", len(by_inn), len(inns))
                if idx < len(pending) and gap > 0:
                    await asyncio.sleep(gap)
        finally:
            await page.close()

    counts: dict[str, int] = {}
    for item in by_inn.values():
        light = item.get("Светофор") or "Нет данных"
        counts[light] = counts.get(light, 0) + 1
    logger.info("Итог: %s", counts)


def main() -> None:
    setup_logging()
    log_path = ROOT / "output" / "glavsnab_checko" / "run.log"
    log_path.parent.mkdir(parents=True, exist_ok=True)
    file_handler = logging.FileHandler(log_path, encoding="utf-8")
    file_handler.setFormatter(logging.Formatter("%(asctime)s [%(levelname)s] %(message)s", "%H:%M:%S"))
    logging.getLogger().addHandler(file_handler)
    from cloakbrowser.download import ensure_binary

    logger.info("CloakBrowser: %s", ensure_binary())
    parser = argparse.ArgumentParser(description="Скоринг базы Главснаб через Checko")
    parser.add_argument("--xlsx", default=str(DEFAULT_XLSX))
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--delay", type=float, default=10.0)
    parser.add_argument("--gap", type=float, default=8.0, help="Пауза между компаниями, секунды")
    parser.add_argument("--headless", action="store_true", default=True)
    args = parser.parse_args()
    asyncio.run(
        run(
            Path(args.xlsx),
            limit=args.limit or None,
            delay=args.delay,
            gap=args.gap,
            headless=args.headless,
        )
    )


if __name__ == "__main__":
    main()
